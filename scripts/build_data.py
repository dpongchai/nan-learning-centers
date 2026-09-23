"""แปลงไฟล์ Excel (data/Nan_learning_centers.xlsx) เป็น data/centers.json สำหรับเว็บไซต์

วิธีใช้:  python scripts/build_data.py
ต้องมี:   pip install openpyxl

สคริปต์จะตรวจข้อมูลและแสดงคำเตือน เช่น พิกัดอยู่นอกจังหวัดน่าน
หรือพิกัดไม่ตรงกับอำเภอ/ตำบลที่กรอกไว้ แต่จะยังสร้างไฟล์ให้ตามปกติ
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "Nan_learning_centers.xlsx"
OUT = ROOT / "data" / "centers.json"
GEO = ROOT / "data" / "nan-boundaries.json"
SHEET = "LearningCenter"

# สีหมุดของแต่ละหมวด (theme_key ในไฟล์ Excel) เพิ่มหมวดใหม่ได้ที่นี่
THEME_COLORS = {
    "farm": "#6FA84A",
    "forest": "#2F8F7A",
    "bee": "#E3B23C",
    "coffee": "#A0694B",
    "community": "#D9678E",
    "academic": "#6C8FD6",
    "craft": "#B07CC6",
}
THEME_EN = {
    "farm": "Natural & organic farming",
    "forest": "Community forests",
    "bee": "Bees & honey",
    "coffee": "Coffee",
    "community": "Community development",
    "academic": "University & research",
    "craft": "Crafts & community business",
}
COURSE = {"มี": "yes", "ไม่มี": "no", "กำลังทำ": "planned", "ไม่ทราบ": "unknown"}

# ชื่อหัวคอลัมน์ในไฟล์ Excel -> ชื่อฟิลด์ใน JSON (ลำดับคอลัมน์สลับได้)
COLUMNS = {
    "Code": "id", "Name_TH": "th", "Name_EN": "en", "Category": "category", "theme_key": "theme",
    "District_TH": "amp", "Sub-district_TH": "tam", "District_EN": "amp_en", "Sub-district_EN": "tam_en",
    "lat": "lat", "lng": "lng", "Link": "map_url", "Description_TH": "d_th", "Description_EN": "d_en",
    "มีหลักสูตร": "course", "ชื่อหลักสูตร": "course_name", "เบอร์โทร": "phone", "LINE / Facebook": "social",
    "เวลาเปิด": "hours", "ต้องจองก่อน": "booking", "ค่าใช้จ่าย": "fee", "ลิงก์รูปภาพ": "photo",
    "หมายเหตุการเข้าชม": "visit_note", "Note": "note",
}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, str):
        v = v.replace("\\n", "\n")
        v = re.sub(r"[ \t]+", " ", v)
        return "\n".join(line.strip() for line in v.strip().splitlines()).strip()
    return v


def point_in_ring(x, y, ring):
    inside = False
    j = len(ring) - 1
    for i in range(len(ring)):
        xi, yi = ring[i]
        xj, yj = ring[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def point_in_geom(x, y, geom):
    polys = [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]
    for poly in polys:
        if point_in_ring(x, y, poly[0]) and not any(point_in_ring(x, y, h) for h in poly[1:]):
            return True
    return False


def main():
    ws = load_workbook(XLSX, data_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]
    idx = {COLUMNS[h]: i for i, h in enumerate(header) if h in COLUMNS}
    missing = [c for c in ("id", "th", "theme", "lat", "lng") if c not in idx]
    if missing:
        sys.exit(f"ไม่พบคอลัมน์ที่จำเป็น: {missing}")

    tambons = json.loads(GEO.read_text(encoding="utf-8"))["tam"]["features"] if GEO.exists() else []
    warnings, centers, themes, seen = [], [], {}, set()

    for n, row in enumerate(rows[1:], start=2):
        r = {k: clean(row[i]) if i < len(row) else "" for k, i in idx.items()}
        if not r.get("id") or not r.get("th"):
            continue
        cid = r["id"]
        if cid in seen:
            warnings.append(f"แถว {n}: รหัส {cid} ซ้ำ")
        seen.add(cid)

        theme = r.get("theme") or "other"
        if theme not in THEME_COLORS:
            warnings.append(f"{cid}: theme_key '{theme}' ไม่รู้จัก (จะใช้สีเทา)")
        themes.setdefault(theme, [r.get("category") or theme, THEME_EN.get(theme, r.get("category") or theme),
                                  THEME_COLORS.get(theme, "#8A94A6")])

        try:
            lat = float(r["lat"]) if r.get("lat") != "" else None
            lng = float(r["lng"]) if r.get("lng") != "" else None
        except (TypeError, ValueError):
            lat = lng = None
            warnings.append(f"{cid}: lat/lng ไม่ใช่ตัวเลข")
        if lat is None or lng is None:
            warnings.append(f"{cid}: ยังไม่มีพิกัด จะแสดงในรายการ 'ยังไม่มีที่ตั้ง'")
        elif tambons:
            hit = next((f["properties"] for f in tambons if point_in_geom(lng, lat, f["geometry"])), None)
            if not hit:
                warnings.append(f"{cid}: พิกัด {lat:.5f},{lng:.5f} อยู่นอกจังหวัดน่าน")
            elif (hit["amp"], hit["th"]) != (r.get("amp"), r.get("tam")):
                warnings.append(f"{cid}: พิกัดอยู่ใน ต.{hit['th']} อ.{hit['amp']} แต่ในไฟล์กรอก "
                                f"ต.{r.get('tam')} อ.{r.get('amp')}")

        d_th = r.get("d_th") or r.get("note") or ""
        course_name = r.get("course_name") or ""
        more_url = ""
        if isinstance(course_name, str) and course_name.startswith("http"):
            more_url, course_name = course_name, ""

        centers.append({
            "id": cid, "th": r["th"], "en": r.get("en") or r["th"], "theme": theme,
            "amp": r.get("amp", ""), "tam": r.get("tam", ""), "amp_en": r.get("amp_en", ""), "tam_en": r.get("tam_en", ""),
            "lat": round(lat, 6) if lat is not None else None, "lng": round(lng, 6) if lng is not None else None,
            "map_url": r.get("map_url", ""), "d_th": d_th, "d_en": r.get("d_en", ""),
            "course": COURSE.get(r.get("course"), "unknown"), "course_name": course_name, "more_url": more_url,
            "phone": str(r.get("phone", "")), "social": r.get("social", ""), "hours": r.get("hours", ""),
            "booking": r.get("booking", ""), "fee": str(r.get("fee", "")), "photo": r.get("photo", ""),
            "visit_note": r.get("visit_note", ""),
        })

    ordered = {k: themes[k] for k in THEME_COLORS if k in themes}
    ordered.update({k: v for k, v in themes.items() if k not in ordered})
    OUT.write_text(json.dumps({"themes": ordered, "centers": centers}, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"สร้าง {OUT.relative_to(ROOT)} แล้ว: {len(centers)} ศูนย์")
    if warnings:
        print(f"\nคำเตือน {len(warnings)} รายการ (ไม่หยุดการสร้างไฟล์):")
        for w in warnings:
            print("  -", w)


if __name__ == "__main__":
    main()
